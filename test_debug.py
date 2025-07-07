#!/usr/bin/env python3

import sys
import os
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def test_cell_counting(excel_file):
    """Test the exact same logic as the app to see what's happening"""
    print(f"Testing {excel_file}...")
    
    try:
        workbook = load_workbook(excel_file, data_only=True)
        worksheet = workbook.active
        print(f"Worksheet: {worksheet.title}")
        print(f"Max row: {worksheet.max_row}, Max col: {worksheet.max_column}")
        
        # Get bounds of worksheet but limit for performance
        max_search_row = min(worksheet.max_row, 1000)
        max_search_col = min(worksheet.max_column, 100)
        print(f"Search bounds: {max_search_row} rows, {max_search_col} columns")
        
        visible_data_cells = []
        hidden_rows_found = 0
        hidden_cols_found = 0
        
        # Debug: Check for hidden rows and columns
        print("Checking for hidden rows and columns...")
        for row_num in range(1, max_search_row + 1):
            if worksheet.row_dimensions[row_num].hidden:
                hidden_rows_found += 1
                print(f"  Row {row_num} is hidden")
        
        for col_num in range(1, max_search_col + 1):
            col_letter = get_column_letter(col_num)
            if worksheet.column_dimensions[col_letter].hidden:
                hidden_cols_found += 1
                print(f"  Column {col_letter} is hidden")
        
        print(f"Found {hidden_rows_found} hidden rows and {hidden_cols_found} hidden columns")
        
        # Count visible cells
        for row_num in range(1, max_search_row + 1):
            # Skip hidden rows entirely - simple check like main branch
            if worksheet.row_dimensions[row_num].hidden:
                continue
                
            for col_num in range(1, max_search_col + 1):
                # Skip hidden columns entirely - simple check like main branch
                if worksheet.column_dimensions[get_column_letter(col_num)].hidden:
                    continue
                    
                cell = worksheet.cell(row=row_num, column=col_num)
                
                # Check if cell has data or formatting
                has_data = cell.value is not None
                has_formatting = False
                
                if not has_data:
                    # Quick formatting check
                    if (cell.fill and hasattr(cell.fill, 'fgColor') and cell.fill.fgColor):
                        if ((hasattr(cell.fill.fgColor, 'rgb') and cell.fill.fgColor.rgb and 
                             cell.fill.fgColor.rgb not in ['FFFFFFFF', '00000000']) or
                            (hasattr(cell.fill.fgColor, 'indexed') and cell.fill.fgColor.indexed is not None and 
                             cell.fill.fgColor.indexed not in [64, 0])):
                            has_formatting = True
                    
                    if not has_formatting and cell.border:
                        has_formatting = any([
                            cell.border.top and cell.border.top.style,
                            cell.border.bottom and cell.border.bottom.style,
                            cell.border.left and cell.border.left.style,
                            cell.border.right and cell.border.right.style
                        ])
                
                if has_data or has_formatting:
                    visible_data_cells.append((row_num, col_num))
                    print(f"  Found visible cell at {row_num},{col_num}: {cell.value}")
        
        print(f"Found {len(visible_data_cells)} visible cells with data/formatting")
        
        if visible_data_cells:
            # Get unique visible rows and columns that have data
            visible_rows_set = set(cell[0] for cell in visible_data_cells)
            visible_cols_set = set(cell[1] for cell in visible_data_cells)
            
            visible_rows = sorted(visible_rows_set)
            visible_cols = sorted(visible_cols_set)
            
            print(f"Unique visible rows: {visible_rows}")
            print(f"Unique visible columns: {visible_cols}")
            print(f"Visible data spans: {len(visible_rows)} rows, {len(visible_cols)} columns")
        
        # Now let's check ALL cells to see if there are any that might be incorrectly counted
        print("\nChecking ALL cells for comparison:")
        all_cells_with_data = []
        for row_num in range(1, max_search_row + 1):
            for col_num in range(1, max_search_col + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                if cell.value is not None:
                    all_cells_with_data.append((row_num, col_num))
                    is_hidden_row = worksheet.row_dimensions[row_num].hidden
                    is_hidden_col = worksheet.column_dimensions[get_column_letter(col_num)].hidden
                    status = "HIDDEN" if (is_hidden_row or is_hidden_col) else "VISIBLE"
                    print(f"  Cell {row_num},{col_num}: {cell.value} - {status}")
        
        if all_cells_with_data:
            all_rows = sorted(set(cell[0] for cell in all_cells_with_data))
            all_cols = sorted(set(cell[1] for cell in all_cells_with_data))
            print(f"ALL data (including hidden): {len(all_rows)} rows, {len(all_cols)} columns")
        
        workbook.close()
        
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    if len(sys.argv) > 1:
        test_cell_counting(sys.argv[1])
    else:
        print("Usage: python test_debug.py <excel_file>")
        print("Available files:")
        for f in ["test_hidden.xlsx", "comprehensive_test.xlsx"]:
            if os.path.exists(f):
                print(f"  {f}")