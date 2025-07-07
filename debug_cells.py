#!/usr/bin/env python3

import sys
import os
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def debug_visible_cells(excel_file):
    """Debug function to check visible cell counting"""
    print(f"Loading {excel_file}...")
    
    try:
        workbook = load_workbook(excel_file, data_only=True)
        worksheet = workbook.active
        print(f"Worksheet: {worksheet.title}")
        
        # Get bounds of worksheet but limit for performance
        max_search_row = min(worksheet.max_row, 1000)
        max_search_col = min(worksheet.max_column, 100)
        
        print(f"Max search bounds: {max_search_row} rows, {max_search_col} columns")
        
        visible_data_cells = []
        hidden_rows_found = 0
        hidden_cols_found = 0
        
        # Check for hidden rows and columns
        print("Checking for hidden rows and columns...")
        for row_num in range(1, max_search_row + 1):
            row_dim = worksheet.row_dimensions[row_num]
            is_row_hidden = (hasattr(row_dim, 'hidden') and row_dim.hidden) or (hasattr(row_dim, 'height') and row_dim.height == 0)
            if is_row_hidden:
                hidden_rows_found += 1
        
        for col_num in range(1, max_search_col + 1):
            col_letter = get_column_letter(col_num)
            col_dim = worksheet.column_dimensions[col_letter]
            is_col_hidden = (hasattr(col_dim, 'hidden') and col_dim.hidden) or (hasattr(col_dim, 'width') and col_dim.width == 0)
            if is_col_hidden:
                hidden_cols_found += 1
        
        print(f"Found {hidden_rows_found} hidden rows and {hidden_cols_found} hidden columns")
        
        # Now count visible cells with data
        for row_num in range(1, max_search_row + 1):
            row_dim = worksheet.row_dimensions[row_num]
            is_row_hidden = (hasattr(row_dim, 'hidden') and row_dim.hidden) or (hasattr(row_dim, 'height') and row_dim.height == 0)
            
            if is_row_hidden:
                continue
                
            for col_num in range(1, max_search_col + 1):
                col_letter = get_column_letter(col_num)
                col_dim = worksheet.column_dimensions[col_letter]
                is_col_hidden = (hasattr(col_dim, 'hidden') and col_dim.hidden) or (hasattr(col_dim, 'width') and col_dim.width == 0)
                
                if is_col_hidden:
                    continue
                    
                cell = worksheet.cell(row=row_num, column=col_num)
                
                # Check if cell has data or formatting
                has_data = cell.value is not None
                has_formatting = False
                
                if not has_data:
                    # Quick formatting check
                    if (cell.fill and hasattr(cell.fill, 'fgColor') and cell.fill.fgColor):
                        if ((hasattr(cell.fill.fgColor, 'rgb') and cell.fill.fgColor.rgb and 
                             cell.fill.fgColor.rgb not in ['FFFFFFFF', '00000000']) or
                            (hasattr(cell.fill.fgColor, 'indexed') and cell.fill.fgColor.indexed is not None and 
                             cell.fill.fgColor.indexed not in [64, 0])):
                            has_formatting = True
                    
                    if not has_formatting and cell.border:
                        has_formatting = any([
                            cell.border.top and cell.border.top.style,
                            cell.border.bottom and cell.border.bottom.style,
                            cell.border.left and cell.border.left.style,
                            cell.border.right and cell.border.right.style
                        ])
                
                if has_data or has_formatting:
                    visible_data_cells.append((row_num, col_num))
        
        print(f"Found {len(visible_data_cells)} visible cells with data/formatting")
        
        if visible_data_cells:
            # Get unique visible rows and columns that have data
            visible_rows_set = set(cell[0] for cell in visible_data_cells)
            visible_cols_set = set(cell[1] for cell in visible_data_cells)
            
            print(f"Unique visible rows: {len(visible_rows_set)}")
            print(f"Unique visible columns: {len(visible_cols_set)}")
            print(f"Visible data spans: {len(visible_rows_set)} rows, {len(visible_cols_set)} columns")
            
            # Show some sample cells
            print("\nSample visible cells (first 10):")
            for i, (row, col) in enumerate(visible_data_cells[:10]):
                cell = worksheet.cell(row=row, column=col)
                print(f"  Row {row}, Col {col}: {cell.value}")
        
        workbook.close()
        
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        debug_visible_cells(sys.argv[1])
    else:
        print("Usage: python debug_cells.py <excel_file>")